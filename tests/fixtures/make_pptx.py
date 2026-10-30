"""Builds tests/fixtures/sample.pptx: one slide per PowerPoint feature the importer handles.
Run with: python3 tests/fixtures/make_pptx.py   (needs python-pptx and Pillow)"""
import io, os
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE
from lxml import etree
from PIL import Image, ImageDraw

HERE = os.path.dirname(os.path.abspath(__file__))
prs = Presentation(); prs.slide_width = Inches(13.333); prs.slide_height = Inches(7.5)
L = prs.slide_layouts

# 1 title slide
s = prs.slides.add_slide(L[0]); s.shapes.title.text = 'Neural networks from scratch'; s.placeholders[1].text = 'AI 4200 · Week 3'

# 2 bullets, runs, hyperlink, notes, build by paragraph
s = prs.slides.add_slide(L[1]); s.shapes.title.text = 'What a layer computes'
tf = s.placeholders[1].text_frame; tf.text = 'Multiply inputs by weights'
p = tf.add_paragraph(); r = p.add_run(); r.text = 'Add a '; r2 = p.add_run(); r2.text = 'bias'; r2.font.bold = True; r2.font.color.rgb = RGBColor(0xE0, 0x40, 0x7A)
p = tf.add_paragraph(); p.text = 'One output per neuron'; p.level = 1
p = tf.add_paragraph(); r = p.add_run(); r.text = 'Read more'; r.hyperlink.address = 'https://pytorch.org/docs/stable/nn.html'
s.notes_slide.notes_text_frame.text = 'Ask the class what the bias does.'
body_id = s.placeholders[1].shape_id

# 3 shapes, group, connector, rotation
s = prs.slides.add_slide(L[5]); s.shapes.title.text = 'Shapes'
a = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1), Inches(2), Inches(3), Inches(1.5)); a.text = 'Input'; a.fill.solid(); a.fill.fore_color.rgb = RGBColor(0xDC, 0xE5, 0xFA); a.line.color.rgb = RGBColor(0x25, 0x57, 0xD6)
b = s.shapes.add_shape(MSO_SHAPE.OVAL, Inches(8), Inches(2), Inches(2), Inches(2)); b.text = 'σ'
c = s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(4), Inches(2.75), Inches(8), Inches(3)); c.line.width = Pt(2)
ln = c.line._get_or_add_ln(); ln.append(etree.fromstring('<a:tailEnd xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" type="triangle"/>'))
t = s.shapes.add_shape(MSO_SHAPE.ISOSCELES_TRIANGLE, Inches(10.5), Inches(5), Inches(1.5), Inches(1.2)); t.rotation = 30
g = s.shapes.add_group_shape(); g1 = g.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(1), Inches(5), Inches(1), Inches(1)); g2 = g.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(2.5), Inches(5), Inches(1), Inches(1)); g2.text = 'G'
fly_id = b.shape_id

# 4 pictures (alt text, crop)
s = prs.slides.add_slide(L[5]); s.shapes.title.text = 'Pictures'
im = Image.new('RGB', (400, 200), '#2557D6'); ImageDraw.Draw(im).rectangle([0, 0, 199, 199], fill='#E0407A'); buf = io.BytesIO(); im.save(buf, 'PNG')
pic = s.shapes.add_picture(io.BytesIO(buf.getvalue()), Inches(1), Inches(2), Inches(4), Inches(2)); pic._element.nvPicPr.cNvPr.set('descr', 'Pink and blue halves')
pic2 = s.shapes.add_picture(io.BytesIO(buf.getvalue()), Inches(7), Inches(2), Inches(2), Inches(2)); pic2.crop_right = 0.5; pic2._element.nvPicPr.cNvPr.set('descr', 'Pink half only')

# 5 table
s = prs.slides.add_slide(L[5]); s.shapes.title.text = 'Results'
tb = s.shapes.add_table(3, 3, Inches(1), Inches(2), Inches(8), Inches(1.5)).table
for i, row in enumerate([['Model', 'Params', 'Accuracy'], ['Small', '10M', '81%'], ['Large', '340M', '91%']]):
    for j, v in enumerate(row): tb.cell(i, j).text = v

# 6 chart
s = prs.slides.add_slide(L[5]); s.shapes.title.text = 'A chart'
cd = CategoryChartData(); cd.categories = ['2019', '2020', '2021']; cd.add_series('Accuracy', (61, 78, 94))
s.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(1), Inches(1.8), Inches(8), Inches(4.5), cd)

# 7 hidden slide
s = prs.slides.add_slide(L[5]); s.shapes.title.text = 'Backup slide'; s._element.set('show', '0')

# animations and a transition (python-pptx has no API for these)
P = 'http://schemas.openxmlformats.org/presentationml/2006/main'
def timing(sl, spid, preset, subtype, cls='entr', by_para=False, paras=1):
    pars = ''
    for k in range(paras):
        tgt = f'<p:spTgt spid="{spid}">' + (f'<p:txEl><p:pRg st="{k}" end="{k}"/></p:txEl>' if by_para else '') + '</p:spTgt>'
        pars += f'''<p:par><p:cTn id="{10+k*4}" fill="hold"><p:stCondLst><p:cond delay="indefinite"/></p:stCondLst><p:childTnLst><p:par><p:cTn id="{11+k*4}" fill="hold"><p:stCondLst><p:cond delay="0"/></p:stCondLst><p:childTnLst>
<p:par><p:cTn id="{12+k*4}" presetID="{preset}" presetClass="{cls}" presetSubtype="{subtype}" fill="hold" nodeType="clickEffect"><p:stCondLst><p:cond delay="0"/></p:stCondLst><p:childTnLst>
<p:set><p:cBhvr><p:cTn id="{13+k*4}" dur="1" fill="hold"><p:stCondLst><p:cond delay="0"/></p:stCondLst></p:cTn><p:tgtEl>{tgt}</p:tgtEl><p:attrNameLst><p:attrName>style.visibility</p:attrName></p:attrNameLst></p:cBhvr><p:to><p:strVal val="visible"/></p:to></p:set>
</p:childTnLst></p:cTn></p:par></p:childTnLst></p:cTn></p:par></p:childTnLst></p:cTn></p:par>'''
    xml = f'''<p:timing xmlns:p="{P}"><p:tnLst><p:par><p:cTn id="1" dur="indefinite" restart="never" nodeType="tmRoot"><p:childTnLst><p:seq concurrent="1" nextAc="seek"><p:cTn id="2" dur="indefinite" nodeType="mainSeq"><p:childTnLst>{pars}</p:childTnLst></p:cTn>
<p:prevCondLst><p:cond evt="onPrev" delay="0"><p:tgtEl><p:sldTgt/></p:tgtEl></p:cond></p:prevCondLst><p:nextCondLst><p:cond evt="onNext" delay="0"><p:tgtEl><p:sldTgt/></p:tgtEl></p:cond></p:nextCondLst></p:seq></p:childTnLst></p:cTn></p:par></p:tnLst></p:timing>'''
    sl._element.append(etree.fromstring(xml))
timing(prs.slides[1], body_id, 10, 0, by_para=True, paras=4)
timing(prs.slides[2], fly_id, 2, 4)
prs.slides[2]._element.insert(2, etree.fromstring(f'<p:transition xmlns:p="{P}" spd="slow"><p:push dir="u"/></p:transition>'))
out = os.path.join(HERE, 'sample.pptx'); prs.save(out); print('wrote', out)
